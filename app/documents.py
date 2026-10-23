from __future__ import annotations

from contextlib import closing
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import base64
import shutil

from .database import connect, database_path, initialize
from .retrieval import index_document


SUPPORTED_SUFFIXES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
MAX_TEXT_FILE_SIZE = 2 * 1024 * 1024
MAX_PRESENTATION_SIZE = 25 * 1024 * 1024


@dataclass(frozen=True)
class SourceDocument:
    id: int
    course_id: int
    filename: str
    media_type: str
    size_bytes: int
    content_hash: str
    status: str
    created_at: str

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


class CourseNotFoundError(LookupError):
    pass


class DocumentNotFoundError(LookupError):
    pass


class DuplicateDocumentError(ValueError):
    pass


class DocumentValidationError(ValueError):
    pass


class DocumentRepository:
    def __init__(
        self,
        database: Path | None = None,
        storage_root: Path | None = None,
    ) -> None:
        self.database = database
        initialize(database)
        selected_database = database or database_path()
        self.storage_root = storage_root or selected_database.parent / "uploads"
        self.storage_root.mkdir(parents=True, exist_ok=True)

    def create(
        self,
        course_id: int,
        filename: str,
        content: str = "",
        file_data_url: str = "",
    ) -> SourceDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise DocumentValidationError("目前只支持 .txt、.md 和 .pptx 文件")

        if suffix == ".pptx":
            encoded, index_content = self._read_presentation(file_data_url)
            max_size = MAX_PRESENTATION_SIZE
        else:
            try:
                encoded = content.encode("utf-8")
            except UnicodeEncodeError as error:
                raise DocumentValidationError("文件必须是 UTF-8 文本") from error
            if not content.strip():
                raise DocumentValidationError("文件内容不能为空")
            index_content = content
            max_size = MAX_TEXT_FILE_SIZE
        if len(encoded) > max_size:
            limit_mb = max_size // (1024 * 1024)
            raise DocumentValidationError(f"文件不能超过 {limit_mb} MB")

        content_hash = sha256(encoded).hexdigest()
        created_at = datetime.now(timezone.utc).isoformat()
        media_type = SUPPORTED_SUFFIXES[suffix]

        with closing(connect(self.database)) as connection:
            course = connection.execute(
                "SELECT id FROM courses WHERE id = ?", (course_id,)
            ).fetchone()
            if course is None:
                raise CourseNotFoundError(f"课程 {course_id} 不存在")
            duplicate = connection.execute(
                "SELECT id FROM source_documents WHERE course_id = ? AND content_hash = ?",
                (course_id, content_hash),
            ).fetchone()
            if duplicate is not None:
                raise DuplicateDocumentError("这份资料已经上传过")

            course_directory = self.storage_root / str(course_id)
            course_directory.mkdir(parents=True, exist_ok=True)
            safe_name = f"{content_hash[:16]}{suffix}"
            stored_file = course_directory / safe_name
            stored_file.write_bytes(encoded)

            try:
                cursor = connection.execute(
                    """
                    INSERT INTO source_documents (
                        course_id, filename, media_type, size_bytes, content_hash,
                        storage_path, status, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, 'ready', ?)
                    """,
                    (
                        course_id,
                        Path(filename).name,
                        media_type,
                        len(encoded),
                        content_hash,
                        str(stored_file),
                        created_at,
                    ),
                )
                index_document(
                    connection,
                    int(cursor.lastrowid),
                    course_id,
                    index_content,
                    "text/markdown" if suffix == ".pptx" else media_type,
                )
                connection.commit()
            except Exception:
                stored_file.unlink(missing_ok=True)
                raise

        return SourceDocument(
            int(cursor.lastrowid),
            course_id,
            Path(filename).name,
            media_type,
            len(encoded),
            content_hash,
            "ready",
            created_at,
        )

    @staticmethod
    def _read_presentation(file_data_url: str) -> tuple[bytes, str]:
        prefix = (
            "data:application/vnd.openxmlformats-officedocument."
            "presentationml.presentation;base64,"
        )
        if not file_data_url.startswith(prefix):
            raise DocumentValidationError("PPTX 文件数据无效")
        try:
            encoded = base64.b64decode(file_data_url[len(prefix):], validate=True)
            from pptx import Presentation

            presentation = Presentation(BytesIO(encoded))
        except Exception as error:
            raise DocumentValidationError("无法读取 PPTX，请确认文件没有损坏") from error
        slides: list[str] = []
        for number, slide in enumerate(presentation.slides, 1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            title = texts[0].splitlines()[0] if texts else f"幻灯片 {number}"
            body = "\n\n".join(texts)
            if body:
                slides.append(f"# 第 {number} 张：{title}\n\n{body}")
        extracted = "\n\n".join(slides).strip()
        if not extracted:
            raise DocumentValidationError("PPTX 中没有找到可检索的文字")
        return encoded, extracted

    def list_for_course(self, course_id: int) -> list[SourceDocument]:
        with closing(connect(self.database)) as connection:
            course = connection.execute(
                "SELECT id FROM courses WHERE id = ?", (course_id,)
            ).fetchone()
            if course is None:
                raise CourseNotFoundError(f"课程 {course_id} 不存在")
            rows = connection.execute(
                """
                SELECT id, course_id, filename, media_type, size_bytes,
                       content_hash, status, created_at
                FROM source_documents
                WHERE course_id = ?
                ORDER BY id DESC
                """,
                (course_id,),
            ).fetchall()
        return [SourceDocument(**dict(row)) for row in rows]

    def delete(self, course_id: int, document_id: int) -> None:
        with closing(connect(self.database)) as connection:
            row = connection.execute(
                """
                SELECT storage_path FROM source_documents
                WHERE id = ? AND course_id = ?
                """,
                (document_id, course_id),
            ).fetchone()
            if row is None:
                raise DocumentNotFoundError(f"资料 {document_id} 不存在")
            connection.execute(
                "DELETE FROM source_documents WHERE id = ? AND course_id = ?",
                (document_id, course_id),
            )
            connection.commit()
        Path(row["storage_path"]).unlink(missing_ok=True)

    def delete_course_storage(self, course_id: int) -> None:
        shutil.rmtree(self.storage_root / str(course_id), ignore_errors=True)

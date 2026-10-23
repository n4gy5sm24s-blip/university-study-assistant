from pathlib import Path
from io import BytesIO
import base64
import shutil
import unittest
from uuid import uuid4
from pptx import Presentation

from app.courses import CourseRepository
from app.documents import (
    CourseNotFoundError,
    DocumentRepository,
    DocumentValidationError,
    DuplicateDocumentError,
    SUPPORTED_SUFFIXES,
)


class DocumentRepositoryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(__file__).parent / ".tmp" / uuid4().hex
        self.root.mkdir(parents=True)
        self.database = self.root / "test.db"
        self.storage = self.root / "uploads"
        self.courses = CourseRepository(self.database)
        self.documents = DocumentRepository(self.database, self.storage)
        self.course = self.courses.create("测试课程")

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def test_upload_and_list_text_document(self) -> None:
        created = self.documents.create(
            self.course.id, "lesson.md", "# 第一课\n这是自制测试资料。"
        )
        listed = self.documents.list_for_course(self.course.id)
        self.assertEqual(listed, [created])
        self.assertEqual(created.status, "ready")
        self.assertEqual(len(list(self.storage.rglob("*.md"))), 1)

    def test_rejects_unsupported_and_empty_documents(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self.documents.create(self.course.id, "lesson.pdf", "内容")
        with self.assertRaises(DocumentValidationError):
            self.documents.create(self.course.id, "lesson.txt", "   ")

    def test_upload_pptx_extracts_slide_text_for_search(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "心理安全"
        slide.placeholders[1].text = "心理安全让团队成员能够提出问题并承认错误。"
        buffer = BytesIO()
        presentation.save(buffer)
        data_url = (
            "data:application/vnd.openxmlformats-officedocument."
            "presentationml.presentation;base64,"
            + base64.b64encode(buffer.getvalue()).decode("ascii")
        )
        created = self.documents.create(
            self.course.id, "week4.pptx", file_data_url=data_url
        )
        self.assertEqual(created.status, "ready")
        self.assertEqual(created.media_type, SUPPORTED_SUFFIXES[".pptx"])
        self.assertEqual(len(list(self.storage.rglob("*.pptx"))), 1)

        from app.retrieval import RetrievalRepository

        results = RetrievalRepository(self.database).search(
            self.course.id, "心理安全"
        )
        self.assertEqual(results[0].filename, "week4.pptx")
        self.assertEqual(results[0].section_title, "第 1 张：心理安全")

    def test_rejects_duplicate_content_in_same_course(self) -> None:
        self.documents.create(self.course.id, "one.txt", "相同内容")
        with self.assertRaises(DuplicateDocumentError):
            self.documents.create(self.course.id, "two.txt", "相同内容")

    def test_rejects_missing_course(self) -> None:
        with self.assertRaises(CourseNotFoundError):
            self.documents.create(999, "lesson.txt", "测试内容")

    def test_delete_removes_database_record_and_file(self) -> None:
        created = self.documents.create(
            self.course.id, "lesson.txt", "将被删除的测试内容"
        )
        self.documents.delete(self.course.id, created.id)
        self.assertEqual(self.documents.list_for_course(self.course.id), [])
        self.assertEqual(list(self.storage.rglob("*.txt")), [])


if __name__ == "__main__":
    unittest.main()
